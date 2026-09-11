import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.text_extraction import (
    NoExtractableText,
    UnsupportedFileType,
    extract_text_from_upload,
)


def _blank_pdf_bytes() -> bytes:
    """A structurally valid PDF with a page but no text content - the shape of a
    scanned/image-only PDF as far as text extraction is concerned."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_blank_pdf_raises_no_extractable_text():
    """Regression test: a page-marker-only result (every page empty) must not slip
    through as if real text were found - see the fix in _extract_pdf."""
    with pytest.raises(NoExtractableText):
        extract_text_from_upload("scanned.pdf", _blank_pdf_bytes(), "application/pdf", add_page_markers=True)


def test_blank_pdf_raises_no_extractable_text_without_page_markers():
    with pytest.raises(NoExtractableText):
        extract_text_from_upload("scanned.pdf", _blank_pdf_bytes(), "application/pdf", add_page_markers=False)


def test_extract_docx():
    doc = DocxDocument()
    doc.add_paragraph("First paragraph of the document.")
    doc.add_paragraph("Second paragraph with more detail.")
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text_from_upload("memo.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "First paragraph of the document." in text
    assert "Second paragraph with more detail." in text


def test_extract_plain_text():
    text = extract_text_from_upload("notes.txt", b"Some case notes here.", "text/plain")
    assert text == "Some case notes here."


def test_extract_unsupported_type_raises():
    with pytest.raises(UnsupportedFileType):
        extract_text_from_upload("archive.zip", b"not a real zip", "application/zip")
